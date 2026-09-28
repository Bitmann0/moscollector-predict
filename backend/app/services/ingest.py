"""Идемпотентный приём СМВУ и ОДС, сброс дня и история загрузок."""
import csv
import hashlib
import io
import zipfile
from datetime import date, datetime, time
from typing import NamedTuple

import openpyxl
from fastapi import HTTPException
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import and_, delete, or_, select, update
from sqlalchemy.orm import Session

from .. import models
from ..limits import UPLOAD_MAX_BYTES
from ..schemas.common import Page
from ..schemas.events import EventRowIn, IngestBatchOut, OdsRowIn, ResetDayOut
from ..security import CurrentUser
from . import parameters, semantics, settings_store
from .helpers import assume_msk, count, from_db, now_utc, page_of, to_db
from .notifications import publish_safe

# Лимит файла из limits.py, как в routers/ingest.py: роутер читает на байт больше, чтобы
# сервис отличил файл ровно в лимит от файла больше лимита.
MAX_UPLOAD_BYTES = UPLOAD_MAX_BYTES
# Размер пачки для IN (...): SQLite до 3.32 принимает не больше 999 параметров.
LOOKUP_CHUNK = 500
# «Тревожное сообщение» — термин заказчика для записи с флагом «тревожное» (ответ 1,
# analysis/qa_customer_2026-09-28.md).
ALARM_TITLE = "Тревожное сообщение СМВУ"
SERIES_GROUPS = ("fire", "gas")


def _out(row: models.IngestBatch) -> IngestBatchOut:
    return IngestBatchOut(batch_id=row.id, kind=row.kind, received_at=from_db(row.received_at),
                          rows_total=row.rows_total, accepted=row.accepted,
                          duplicates=row.duplicates, rejected=row.rejected,
                          outside_demo_window=row.outside_demo_window, status=row.status)


def _new_batch(db: Session, kind: str, rows_total: int) -> models.IngestBatch:
    row = models.IngestBatch(kind=kind, received_at=now_utc(), rows_total=rows_total,
                             accepted=0, duplicates=0, rejected=0,
                             outside_demo_window=0, status="accepted")
    db.add(row)
    db.flush()
    return row


def _chunks(values: list) -> list[list]:
    return [values[i:i + LOOKUP_CHUNK] for i in range(0, len(values), LOOKUP_CHUNK)]


def _parse_alarm(value: str | bool) -> bool:
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().casefold()
    if normalized in {"t", "true", "1", "да"}:
        return True
    if normalized in {"f", "false", "0", "нет"}:
        return False
    raise ValueError("invalid_alarm")


def _event_from_row(row: EventRowIn) -> models.Event:
    """Событие без класса: класс ставит _save_rows, когда известен тип датчика канала."""
    alarm = _parse_alarm(row.alarm)
    ts = assume_msk(datetime.fromisoformat(f"{row.day.strip()}T{row.time.strip()}"))
    raw = row.value.strip() if row.value else None
    try:
        value_num = float(raw.replace(",", ".")) if raw else None
    except ValueError:
        value_num = None
    identity = f"{row.event_id}|{row.channel_id}|{ts.isoformat()}|{alarm}|{raw or ''}"
    return models.Event(event_id=row.event_id, channel_id=row.channel_id, ts=to_db(ts),
                        alarm=alarm, val_raw=raw, val_num=value_num,
                        row_hash=hashlib.sha256(identity.encode()).hexdigest())


def _existing_hashes(db: Session, hashes: list[str]) -> set[str]:
    found: set[str] = set()
    for chunk in _chunks(hashes):
        found.update(db.scalars(select(models.Event.row_hash)
                                .where(models.Event.row_hash.in_(chunk))))
    return found


class ChannelInfo(NamedTuple):
    sensor_type: str | None
    obj_id: str | None
    complex_id: str | None


def _channels(db: Session, channel_ids: list[int]) -> dict[int, ChannelInfo]:
    """Тип датчика — для класса; объект и комплекс — для ключа серии ППР/ТО."""
    found: dict[int, ChannelInfo] = {}
    for chunk in _chunks(channel_ids):
        rows = db.execute(select(models.RefChannel.id, models.RefChannel.sensor_type,
                                 models.RefChannel.obj_id, models.RefObject.parent_id)
                          .outerjoin(models.RefObject,
                                     models.RefObject.id == models.RefChannel.obj_id)
                          .where(models.RefChannel.id.in_(chunk)))
        found.update((cid, ChannelInfo(st, obj, parent)) for cid, st, obj, parent in rows)
    return found


def _series_event(ref, group: str | None, info: ChannelInfo | None, channel_id: int,
                  ts: datetime) -> semantics.SeriesEvent | None:
    key = semantics.series_key(group, info.obj_id if info else None,
                               info.complex_id if info else None)
    if key is None:
        return None
    return semantics.SeriesEvent(ref=ref, group=group, key=key, channel_id=channel_id, ts=ts)


def mark_series(db: Session, fresh: list[models.Event], channels: dict[int, ChannelInfo],
                rules: semantics.Rules = semantics.DEFAULT_RULES) -> int:
    """Подсказка «вероятно, ППР или ТО: серия…» новым событиям и уже записанным.

    Серию видно, только когда пришло N-е событие, поэтому подсказку получают и
    предыдущие события серии — UPDATE в той же транзакции, что и приём пачки.
    Вызывать до db.add(fresh): запрос видит только прежние строки. Запрос к БД один
    на пачку: события групп серий у тех же объектов и комплексов за
    [первое − 2 окна, последнее + 2 окна] (окно — rules.series_window, проверено 10 мин).
    Этого хватает, чтобы точно пересчитать подсказки в [первое − окно, последнее + окно]
    (semantics.series_hints); дальше новые события ни на одно окно не влияют.
    Возвращает число обновлённых старых строк.
    """
    new = [e for e in (_series_event(event, event.incident_group,
                                     channels.get(event.channel_id), event.channel_id,
                                     from_db(event.ts)) for event in fresh) if e]
    if not new:
        return 0
    first, last = min(e.ts for e in new), max(e.ts for e in new)
    window = rules.series_window
    stored = db.execute(
        select(models.Event.id, models.Event.channel_id, models.Event.ts,
               models.Event.incident_group, models.Event.hint,
               models.RefChannel.obj_id, models.RefObject.parent_id)
        .join(models.RefChannel, models.RefChannel.id == models.Event.channel_id)
        .outerjoin(models.RefObject, models.RefObject.id == models.RefChannel.obj_id)
        .where(models.Event.incident_group.in_(SERIES_GROUPS),
               models.Event.ts >= to_db(first - 2 * window),
               models.Event.ts <= to_db(last + 2 * window),
               or_(and_(models.Event.incident_group == "fire",
                        models.RefChannel.obj_id.in_({e.key for e in new
                                                      if e.group == "fire"})),
                   and_(models.Event.incident_group == "gas",
                        models.RefObject.parent_id.in_({e.key for e in new
                                                        if e.group == "gas"}))))).all()
    candidates = list(new)
    before: dict[int, tuple[datetime, str | None]] = {}
    for row_id, channel_id, ts, group, hint, obj_id, parent_id in stored:
        event = _series_event(row_id, group, ChannelInfo(None, obj_id, parent_id),
                              channel_id, from_db(ts))
        if event:
            before[row_id] = (event.ts, hint)
            candidates.append(event)
    hints = semantics.series_hints(candidates, rules)
    for event in new:
        if event.ref in hints:
            event.ref.hint = hints[event.ref]
    changed = [{"id": row_id, "hint": hints[row_id]} for row_id, (ts, hint) in before.items()
               if row_id in hints and first - window <= ts <= last + window
               and hint != hints[row_id]]
    if changed:
        db.execute(update(models.Event), changed)
    return len(changed)


def _save_rows(db: Session, rows: list[EventRowIn],
               *, rejected: int = 0, notify: bool = True) -> IngestBatchOut:
    """notify=False — загрузка истории (replay.py --bulk, --catch-up): события получают
    класс, но уведомлений и SSE нет — иначе прошлые тревоги пришли бы диспетчеру как новые.

    Пороги, окна и классы уведомлений — из параметров (parameters.current): одно чтение
    кеша на пачку, БД — не чаще раза в parameters.CACHE_TTL_S."""
    batch = _new_batch(db, "smvu", len(rows) + rejected)
    batch.rejected = rejected
    demo_today = settings_store.demo_today(db)
    tuning = parameters.current(db)
    parsed: list[models.Event] = []
    for item in rows:
        try:
            event = _event_from_row(item)
        except (TypeError, ValueError):
            batch.rejected += 1
            continue
        if from_db(event.ts).date() > demo_today:
            batch.outside_demo_window += 1
            continue
        parsed.append(event)
    # Дубли ищем одним запросом на пачку хешей, а не запросом на строку: файл до 200 МБ —
    # это миллионы строк. seen ловит повтор строки внутри той же партии.
    seen = _existing_hashes(db, list({e.row_hash for e in parsed}))
    channels = _channels(db, list({e.channel_id for e in parsed}))
    fresh: list[models.Event] = []
    for event in parsed:
        if event.row_hash in seen:
            batch.duplicates += 1
            continue
        seen.add(event.row_hash)
        info = channels.get(event.channel_id)
        event.event_class, event.hint, event.incident_group = semantics.classify(
            info.sensor_type if info else None, event.val_raw, event.val_num, event.alarm,
            ts=from_db(event.ts), rules=tuning.rules)
        event.batch_id = batch.id
        fresh.append(event)
    mark_series(db, fresh, channels, tuning.rules)
    alarms: list[models.Notification] = []
    for event in fresh:
        db.add(event)
        batch.accepted += 1
        if notify and tuning.notifies(event.event_class, event.incident_group):
            notification = models.Notification(
                ts=event.ts, kind="event.alarm",
                severity="critical" if event.event_class == "critical" else "warning",
                title=ALARM_TITLE, read_by=[],
                payload={"event_id": event.event_id, "channel_id": event.channel_id,
                         "event_class": event.event_class,
                         "incident_group": event.incident_group})
            db.add(notification)
            alarms.append(notification)
    batch.status = "accepted" if not batch.rejected else ("partial" if batch.accepted else "rejected")
    db.commit()
    # Строки уже в БД: оборванный поток SSE не должен превращать приём в 500.
    for notification in alarms:
        publish_safe("event.alarm", {**notification.payload, "notification_id": notification.id},
                     severity=notification.severity, title=ALARM_TITLE)
    return _out(batch)


def _count_rows(filename: str, content: bytes) -> int:
    if filename.lower().endswith(".xlsx") or content[:2] == b"PK":
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = book.worksheets[0]
            rows = sum(1 for r in sheet.iter_rows(values_only=True)
                       if any(v is not None for v in r))
        finally:
            book.close()
    else:
        text = content.decode("utf-8-sig")
        rows = sum(1 for r in csv.reader(io.StringIO(text)) if any(cell.strip() for cell in r))
    return max(rows - 1, 0)  # первая строка — заголовок


def _cell_text(header: str | None, value):
    """Ячейка XLSX → значение для EventRowIn.

    Excel хранит дату, время и числа как типы, а EventRowIn принимает строки, как в CSV.
    Целые числа пишутся без «.0», чтобы хеш строки совпал с той же строкой из CSV.
    """
    if value is None or isinstance(value, (bool, str)):
        return value
    if isinstance(value, datetime):
        return value.strftime("%H:%M:%S") if header == "время" else value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def ingest_rows(db: Session, rows: list[EventRowIn], user: CurrentUser,
                *, notify: bool = True) -> IngestBatchOut:
    return _save_rows(db, rows, notify=notify)


def ingest_file(db: Session, filename: str, content: bytes, user: CurrentUser,
                *, notify: bool = True) -> IngestBatchOut:
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="file_too_large")
    try:
        rows_total = _count_rows(filename, content)
    except (ValueError, KeyError, OSError, csv.Error, zipfile.BadZipFile, InvalidFileException):
        batch = _new_batch(db, "smvu", 0)
        batch.status = "rejected"
        db.commit()
        return _out(batch)
    try:
        if filename.lower().endswith(".xlsx") or content[:2] == b"PK":
            book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                values = list(book.worksheets[0].iter_rows(values_only=True))
            finally:
                book.close()
            header, data = values[0], values[1:]
            records = [{name: _cell_text(name, cell) for name, cell in zip(header, cells)}
                       for cells in data if any(v is not None for v in cells)]
        else:
            records = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        rows = []
        rejected = 0
        for record in records:
            try:
                rows.append(EventRowIn.model_validate(record))
            except (ValueError, TypeError):
                rejected += 1
    except (IndexError, ValueError, TypeError):
        batch = _new_batch(db, "smvu", rows_total)
        batch.rejected = rows_total
        batch.status = "rejected"
        db.commit()
        return _out(batch)
    return _save_rows(db, rows, rejected=rejected, notify=notify)


def reset_day(db: Session, day: date, user: CurrentUser) -> ResetDayOut:
    start = to_db(assume_msk(datetime.combine(day, datetime.min.time())))
    end = to_db(assume_msk(datetime.combine(day.fromordinal(day.toordinal() + 1), datetime.min.time())))
    result = db.execute(delete(models.Event).where(models.Event.ts >= start, models.Event.ts < end))
    batch = _new_batch(db, "reset", 0)
    batch.accepted = result.rowcount or 0
    db.commit()
    return ResetDayOut(day=day, deleted_events=result.rowcount or 0)


def _same(column, value):
    return column.is_(None) if value is None else column == value


def ingest_ods(db: Session, rows: list[OdsRowIn], user: CurrentUser) -> IngestBatchOut:
    """Повтор той же записи (эмулятор после таймаута) считается дублем, а не новой записью."""
    batch = _new_batch(db, "ods", len(rows))
    seen: set[tuple] = set()
    for row in rows:
        ts = to_db(assume_msk(row.ts))
        key = (ts, row.obj_id, row.record_type, row.decision, row.reason)
        stored = db.scalar(select(models.OdsRecord.id).where(
            models.OdsRecord.ts == ts, _same(models.OdsRecord.obj_id, row.obj_id),
            models.OdsRecord.record_type == row.record_type,
            _same(models.OdsRecord.decision, row.decision),
            _same(models.OdsRecord.reason, row.reason)).limit(1))
        if key in seen or stored is not None:
            batch.duplicates += 1
            continue
        seen.add(key)
        db.add(models.OdsRecord(ts=ts, obj_id=row.obj_id, record_type=row.record_type,
                                decision=row.decision, reason=row.reason, batch_id=batch.id))
        batch.accepted += 1
    db.commit()
    return _out(batch)


def list_batches(db: Session, page: int, page_size: int) -> Page[IngestBatchOut]:
    stmt = select(models.IngestBatch)
    total = count(db, stmt)
    rows = db.scalars(stmt.order_by(models.IngestBatch.id.desc())
                      .offset((page - 1) * page_size).limit(page_size))
    return page_of(IngestBatchOut, [_out(r) for r in rows], total, page, page_size)
