"""Normalize two customer maintenance plans without guessing catalog IDs."""
from __future__ import annotations

import argparse
import datetime as dt
import json
from pathlib import Path

from mkl.maintenance import parse_ppr, parse_to, source_sha256


def normalize(ppr_path: Path, to_path: Path, available_from: dt.date) -> dict:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Install the optional schedules extra: pip install '.[schedules]'") from exc

    ppr_book = openpyxl.load_workbook(ppr_path, read_only=True, data_only=True)
    to_book = openpyxl.load_workbook(to_path, read_only=False, data_only=True)
    try:
        ppr_sheet = ppr_book.active
        to_sheet = to_book.active
        ppr = parse_ppr((i, tuple(row)) for i, row in enumerate(
            ppr_sheet.iter_rows(values_only=True), 1))
        equipment, work = parse_to((i, tuple(cell.value for cell in row),
                                    bool(to_sheet.row_dimensions[i].hidden))
                                   for i, row in enumerate(to_sheet.iter_rows(), 1))
    finally:
        ppr_book.close()
        to_book.close()
    if not ppr or not equipment:
        raise ValueError("one of the schedules contains no recognizable records")
    return {
        "schema_version": 1,
        "available_from": available_from.isoformat(),
        "source_metadata": {
            "ppr": {"filename": ppr_path.name, "sha256": source_sha256(ppr_path)},
            "to": {"filename": to_path.name, "sha256": source_sha256(to_path)},
        },
        "join_policy": "manual_mapping_required; source ordinals are not catalog IDs",
        "ppr_objects": ppr,
        "to_equipment": equipment,
        "to_work_months": work,
        "summary": {
            "ppr_objects": len(ppr),
            "ppr_without_dates": sum("no_planned_dates" in r["quality_flags"] for r in ppr),
            "to_objects": len({r["source_object_key"] for r in equipment}),
            "to_equipment_rows": len(equipment),
            "to_work_months": len(work),
            "hidden_equipment_rows": sum(r["hidden_source_row"] for r in equipment),
            "catalog_matches": 0,
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--ppr", type=Path, required=True)
    parser.add_argument("--to", type=Path, required=True)
    parser.add_argument("--available-from", type=dt.date.fromisoformat, required=True)
    parser.add_argument("--output", type=Path, default=Path("data/interim/maintenance_2026.json"))
    args = parser.parse_args()
    payload = normalize(args.ppr, args.to, args.available_from)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n",
                           encoding="utf-8")
    print(json.dumps(payload["summary"], ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
